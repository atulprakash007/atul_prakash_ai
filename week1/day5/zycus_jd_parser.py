import os
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field

# --- Load environment variables and initialize Groq client ---
load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("Where is API KEY ?")

client = Groq(api_key=my_api_key)
model = "llama-3.3-70b-versatile"  # LLM model used for all extraction/scoring calls


# --- Raw job description text (input) ---
job_description = """
Job Description
Hiring Trainee Business Development Executives | Build a Career Where Technology Meets Business
Most engineering graduates build products. At Zycus, you'll learn how to build global businesses.

If you enjoy solving problems, understanding technology, working with people, and want to accelerate your career beyond coding, this is an opportunity to become part of the next generation of Enterprise SaaS sales leaders. Join one of India's leading AI-first enterprise software companies, recognized as a Leader in the Gartner® Magic Quadrant™ for Source-to-Pay Suites, serving Fortune 1000 companies across North America, Europe, the Middle East and APAC.

This is a 6-month paid internship with structured training, mentorship, and a pathway to a full-time Business Development Associate role based on performance.

Why Choose Business Development at Zycus?
Unlike traditional sales roles, you'll be selling cutting-edge AI software to senior executives at some of the world's largest organizations.

From Day One, you'll interact with:

Chief Procurement Officers (CPOs)

Vice Presidents

Procurement Directors

Global Digital Transformation Leaders

You'll learn how Fortune 1000 companies evaluate technology, make million-dollar buying decisions, and transform their businesses using AI.

This is one of the fastest ways to develop commercial, leadership, and communication skills while gaining global business exposure.

What You'll Do
As a Graduate Business Development Associate, you'll work alongside our global sales teams to create new business opportunities.

You'll:

Research Fortune 1000 companies and identify potential opportunities

Use AI tools, LinkedIn Sales Navigator and digital research to identify decision-makers

Engage senior executives through email, phone, video meetings and social selling

Understand customer business challenges and position Zycus' AI platform

Schedule meetings and product demonstrations for our global sales teams

Build account intelligence and map enterprise buying centres

Analyze industries, competitors and market trends

Maintain high-quality CRM data and sales intelligence

Work closely with Sales, Marketing and Product teams

What You'll Learn During the Internship
This is not just an internship.

It's a structured learning program covering:

Enterprise SaaS Sales

AI & Agentic AI Applications in Business

Consultative Selling

Executive Communication

Social Selling & LinkedIn Prospecting

Account Mapping

Global Business Etiquette

Negotiation Fundamentals

Sales Technology & CRM

Market Research & Competitive Intelligence

By the end of the internship, you'll have skills that are valuable across Sales, Product Management, Customer Success, Consulting and Business Strategy.

Who We're Looking For
We're looking for curious, ambitious engineers who enjoy solving business problems.

Education
B.E. / B.Tech ( 2027 Graduates)

Any engineering discipline

You Should Have
Excellent spoken and written English

Strong analytical and problem-solving ability

Curiosity about technology and business

Confidence interacting with people

Ability to learn quickly

High energy and self-motivation

Strong research and data analysis skills

Competitive mindset and willingness to achieve ambitious goals

No prior sales experience is required. We will train you.
Why Engineering Graduates Succeed Here
Engineers naturally excel at:

Breaking down complex problems

Learning new technologies quickly

Thinking analytically

Understanding enterprise software

Communicating technical concepts simply

If you enjoy technology but also want to work with customers, influence business decisions, and build leadership skills, Business Development offers an exciting alternative career path.

Why Join Zycus?
Join an AI-first global SaaS company transforming procurement for Fortune 1000 enterprises

Work with customers across North America, Europe, Middle East and APAC

Learn directly from experienced global sales leaders

Structured learning, mentorship and career development

Fast-track growth into Business Development, Enterprise Sales, Customer Success or Product roles

High-performance culture with merit-based career progression

Attractive stipend during internship and performance-based full-time offer

Uncapped incentive opportunities after conversion to full-time

 
"""


# --- Pydantic schema for structured job description extraction ---
class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]

jobd_schema = JobD.model_json_schema()

# --- Prompts to extract structured JD data via LLM ---
system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{jobd_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt = f"""
Analyze the following job description:

{job_description}
"""
message_system = {
    "role": "system",
    "content": system_prompt
}
message_user = {
    "role": "user",
    "content": user_prompt
}
response_format = {
    "type": "json_object"
}


messages = [message_system, message_user]

# --- LLM call: extract structured job description ---
response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)

answer = response.choices[0].message.content

raw_json = answer
# print(raw_json)


import json
job_data = json.loads(raw_json)

job = JobD(**job_data)  # validated structured job description

print(job.minimum_experience)
print(job.education_requirements)


# --- Schemas for resume parsing and match scoring ---
class MatchResult(BaseModel):
    score: float
    details: dict

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()

# --- Function: compare a parsed resume against the job description via LLM, return match score ---
def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message = {
        "role": "user",
        "content": prompt
    }
    messages = [message]
    response_format = {
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

# --- Function: parse raw resume text into structured Resume object via LLM ---
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system = {
        "role": "system",
        "content": system_prompt
    }
    message_user = {
        "role": "user",
        "content": user_prompt
    }
    messages = [message_system, message_user]
    response_format = {
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


# --- File reading utilities (PDF / DOCX text extraction) ---
from pypdf import PdfReader
from docx import Document

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

# --- Dispatcher: pick the right reader based on file extension ---
def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


# --- Main pipeline: process every resume in the folder ---

resume_folder = Path("resumes")
all_results = []
for file_path in resume_folder.iterdir():
    
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)                            # llm call1 - parse resume into structured data
    time.sleep(5)                                                        # rate-limit buffer between API calls
    result = final_score(job, parsed_resume)                             # llm caLL2 - score resume against job
   
    time.sleep(5)                                                        # rate-limit buffer between API calls to prevent DOS 
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })

# --- Rank candidates by score, best and worst  ---

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_1 = all_results[-1:]


# --- Output: print top and bottom candidates ---
print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST CANDIDATES")
for candidate in worst_1:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])



    #---------------------------------------------------------Thank You----------------------------------------------------------------